"""Create a fake Long Outage workbook so the pipeline can run without any company data."""
import random
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import config

random.seed(42)

REGIONS = ["Region-01", "Region-02", "Region-03", "Region-04", "Region-05", "Region-06"]
CAUSES = [
    "Mains failure",
    "Battery depleted",
    "Fiber cut",
    "Equipment fault",
    "Access issue",
    "Generator fuel",
]
STATUS = ["Open", "In progress", "Escalated"]
REMARKS = [
    "Field team dispatched",
    "Waiting for grid restoration",
    "Spare part requested",
    "Access permission pending",
    "Fuel refill scheduled",
    "Under observation",
]
COLUMNS = [
    "Site ID", "Region", "Outage Start", "Duration (h)",
    "Root Cause", "Status", "Team", "Remarks",
]


def build_rows(prefix: str, n: int):
    now = datetime(2026, 1, 15, 9, 0)
    rows = []
    for i in range(1, n + 1):
        start = now - timedelta(hours=random.randint(6, 96))
        rows.append([
            f"{prefix}-{1000 + i}",
            random.choice(REGIONS),
            start.strftime("%Y-%m-%d %H:%M"),
            round((now - start).total_seconds() / 3600, 1),
            random.choice(CAUSES),
            random.choice(STATUS),
            f"Team {random.choice('ABCDE')}",
            random.choice(REMARKS),
        ])
    return sorted(rows, key=lambda r: -r[3])


def main():
    config.DATA_DIR.mkdir(exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    for key, (sheet, title) in config.TABLES.items():
        ws = wb.create_sheet(sheet)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
        ws["A1"] = title
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].alignment = Alignment(horizontal="center")
        for c, name in enumerate(COLUMNS, 1):
            cell = ws.cell(row=2, column=c, value=name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center")
        for r, row in enumerate(build_rows(key[-1], 28 if key == "OperatorA" else 22), 3):
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v)
        for col, w in zip("ABCDEFGH", [11, 11, 18, 13, 18, 13, 10, 36]):
            ws.column_dimensions[col].width = w
    wb.save(config.WORKBOOK)
    print(f"Sample workbook written: {config.WORKBOOK}")


if __name__ == "__main__":
    main()
