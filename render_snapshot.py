"""Render each RCA table to a hi-res PNG and a one-page PDF.

Two backends:
  * "excel"      - Windows + Microsoft Excel. Exports the used range to a one-page
                   A3 landscape PDF with ExportAsFixedFormat, rasterises it at 200 DPI
                   with PyMuPDF, then trims the white margin with Pillow. This is the
                   method used in production because it is pixel-identical to Excel.
  * "matplotlib" - Cross-platform fallback used for this demo. Draws the table from
                   the cell values so the pipeline runs anywhere (CI, Linux, macOS).

Hard rule for both: every row and every column must appear. Nothing may be cut off.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageChops

import config


# --------------------------------------------------------------------------- helpers
def autocrop(png_path: Path, pad: int = 20) -> None:
    """Trim the white margin around the rendered table, keep a small padding."""
    img = Image.open(png_path).convert("RGB")
    bg = Image.new("RGB", img.size, (255, 255, 255))
    bbox = ImageChops.difference(img, bg).getbbox()
    if bbox:
        l, t, r, b = bbox
        img = img.crop((max(l - pad, 0), max(t - pad, 0),
                        min(r + pad, img.width), min(b + pad, img.height)))
    img.save(png_path, dpi=(config.DPI, config.DPI))


def read_table(sheet_name: str):
    """Return (title, header, rows) from a sheet laid out as: title row, header row, data."""
    ws = load_workbook(config.WORKBOOK, data_only=True)[sheet_name]
    values = [list(r) for r in ws.iter_rows(values_only=True)]
    title = values[0][0]
    header = [h for h in values[1] if h is not None]
    width = len(header)
    rows = [r[:width] for r in values[2:] if any(v is not None for v in r[:width])]
    return title, header, rows


# --------------------------------------------------------------------------- backends
def render_matplotlib(sheet_name: str, out_png: Path, out_pdf: Path) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    title, header, rows = read_table(sheet_name)
    cell_text = [["" if v is None else str(v) for v in r] for r in rows]

    # Size the figure from the content so nothing is ever clipped.
    col_chars = [max(len(str(h)), *(len(r[i]) for r in cell_text)) + 2 for i, h in enumerate(header)]
    fig_w = max(8.0, sum(col_chars) * 0.085 + 0.6)
    fig_h = 0.9 + 0.26 * (len(cell_text) + 1)

    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    ax.axis("off")
    ax.set_title(title, fontsize=13, fontweight="bold", pad=8)
    table = ax.table(cellText=cell_text, colLabels=header, loc="upper center",
                     cellLoc="center",
                     colWidths=[c / sum(col_chars) for c in col_chars])
    table.auto_set_font_size(False)
    table.set_fontsize(8.5)
    table.scale(1, 1.25)

    for (r, _c), cell in table.get_celld().items():
        cell.set_edgecolor("#B7C4D1")
        if r == 0:
            cell.set_facecolor("#1F4E78")
            cell.set_text_props(color="white", fontweight="bold")
        elif r % 2 == 0:
            cell.set_facecolor("#EEF3F8")

    fig.savefig(out_png, dpi=config.DPI, bbox_inches="tight", facecolor="white")
    fig.savefig(out_pdf, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    autocrop(out_png)


def render_excel(sheet_name: str, out_png: Path, out_pdf: Path) -> None:
    """Production technique (Windows only). Requires: pywin32, pymupdf."""
    import fitz  # PyMuPDF
    import win32com.client as win32

    excel = win32.DispatchEx("Excel.Application")
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Open(str(config.WORKBOOK))
        ws = wb.Worksheets(sheet_name)
        rng = ws.UsedRange
        ps = ws.PageSetup
        ps.PrintArea = rng.Address
        ps.Orientation = 2          # xlLandscape
        ps.PaperSize = 8            # xlPaperA3
        ps.Zoom = False
        ps.FitToPagesWide = 1       # one page wide ...
        ps.FitToPagesTall = 1       # ... and one page tall -> nothing is cut off
        ws.ExportAsFixedFormat(0, str(out_pdf))   # 0 = xlTypePDF
        wb.Close(SaveChanges=False)
    finally:
        excel.Quit()

    doc = fitz.open(out_pdf)
    doc[0].get_pixmap(dpi=config.DPI).save(out_png)
    doc.close()
    autocrop(out_png)
    # Note: CopyPicture / metafile paste was tried first and dropped. Off-screen areas
    # came out black and banded rows went missing. PDF export is reliable.


# --------------------------------------------------------------------------- entry
def render_all(backend: str | None = None) -> list[tuple[str, Path, Path]]:
    backend = backend or ("excel" if sys.platform == "win32" and config.USE_EXCEL else "matplotlib")
    render = render_excel if backend == "excel" else render_matplotlib
    config.OUTPUT_DIR.mkdir(exist_ok=True)
    results = []
    for key, (sheet, _title) in config.TABLES.items():
        png = config.OUTPUT_DIR / f"{key}_Long_Outage_RCA.png"
        pdf = config.OUTPUT_DIR / f"{key}_Long_Outage_RCA.pdf"
        render(sheet, png, pdf)
        print(f"[render:{backend}] {sheet} -> {png.name}, {pdf.name}")
        results.append((key, png, pdf))
    return results


if __name__ == "__main__":
    render_all()
